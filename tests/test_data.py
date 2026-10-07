"""Afmörkuð próf; HR_GAGNAMAPPA virkjar einnig raunverulegan innlestur."""
import copy
import os
from pathlib import Path
import sys
import tempfile
import unittest

import pandas as pd
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import data


def skrifa_daemi(mappa, *, bn=False, duplicate=False, missing=False, unknown_role=False):
    """Smátt tilbúið Excel-dæmi, ekki afrit eða breyting á heimild."""
    w = Workbook()
    s = w.active
    s.title = "Hjúkkur"
    h = ["Nafn", "Vinnufyrirkomulag", *data.HAEFNIDALKAR, "bör", "Vaktir"]
    for m in range(1, 13):
        rest = list(data.MANADARFYRIRSAGNIR)
        if bn and m == 8:
            rest[2] = "BN"
        h.extend([data.MANADARHEITI[m - 1], *rest])
    h.append("Athugasemdir")
    if missing:
        h[0] = "Óþekktur dálkur"
    s.append(h)
    for n, f in (("h1", .525), ("h2", 0)):
        row = [n, f, *([1] * len(data.HAEFNIDALKAR)), 0, " MV - KV - NV "]
        for m in range(1, 13):
            row.extend([f, *([0] * 8)])
        s.append([*row, None])
    if duplicate:
        s.append([c.value for c in s[2]])
    s.append([None, "Samtals stöðugildi"])
    s.append([None, "Munaðarlaus texti"])
    q = w.create_sheet("Sheet2")
    q.append([None, "Mönnunarþörf"])
    q.append([None, "MV", "KV", "NV"])
    q.append([None, " a ", "A", "A"])
    q.append([None, "A", "A", None])
    if unknown_role:
        q.append([None, "Óþekkt", None, None])
    path = Path(mappa) / data.SKRA
    w.save(path)
    return path


class DataTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name)
        skrifa_daemi(self.folder)
        self.t = data.lesa_starfsfolk(self.folder)
        self.settings = dict(fra_drattur=["leyfi"], vaktir_100=20,
                             jafnlangar_vaktir=True, sama_thorf_alla_daga=True)

    def tearDown(self):
        self.tmp.cleanup()

    def test_dates(self):
        for y, m, n in ((2026, 11, 30), (2026, 10, 31), (2028, 2, 29), (2026, 2, 28)):
            self.assertEqual(len(data.bua_til_dagsetningar(y, m)), n)
        for y, m in ((True, 1), (2026, False), (2026, 13), (2026, 1.5), (0, 1)):
            with self.assertRaises(ValueError):
                data.bua_til_dagsetningar(y, m)

    def test_rows_and_duplicate(self):
        self.assertEqual(self.t["Nafn"].tolist(), ["h1", "h2"])
        skrifa_daemi(self.folder, duplicate=True)
        with self.assertRaisesRegex(ValueError, "tvítekin"):
            data.lesa_starfsfolk(self.folder)

    def test_missing_column(self):
        skrifa_daemi(self.folder, missing=True)
        with self.assertRaisesRegex(ValueError, "skyldudálka"):
            data.lesa_starfsfolk(self.folder)

    def test_months(self):
        for m in range(1, 13):
            c = data.velja_manadarblokk(self.t, m)
            self.assertEqual(len(c), 9)
            self.assertEqual(c["hlutfall"], data.MANADARHEITI[m - 1])

    def test_month_order(self):
        t = self.t.copy()
        t.attrs["fyrirsagnir"] = list(self.t.attrs["fyrirsagnir"])
        j = t.attrs["fyrirsagnir"].index("Janúar")
        t.attrs["fyrirsagnir"][j + 1] = "Leyfi"
        with self.assertRaisesRegex(ValueError, "röng dálkaröð"):
            data.velja_manadarblokk(t, 1)

    def test_bn_explicit(self):
        skrifa_daemi(self.folder, bn=True)
        t = data.lesa_starfsfolk(self.folder)
        with self.assertRaisesRegex(ValueError, "BN"):
            data.velja_manadarblokk(t, 8)
        self.assertIn("nm", data.velja_manadarblokk(t, 8, bn_sem_nm=True))

    def test_percentages(self):
        for x, expected in ((.8, .8), ("80%", .8), (" 52,5 % ", .525), ("0,8", .8)):
            t = pd.DataFrame({"Nafn": ["h1"], "p": [x]})
            self.assertAlmostEqual(data.breyta_prosentum(t, ["p"]).iloc[0]["p"], expected)
        for x in (None, "", "abc", 80, float("nan"), float("inf"), True):
            with self.assertRaisesRegex(ValueError, "h1"):
                data.breyta_prosentum(pd.DataFrame({"Nafn": ["h1"], "p": [x]}), ["p"])

    def test_negative_before_filtering(self):
        c = data.velja_manadarblokk(self.t, 11)
        t = self.t.copy()
        t.loc[t["Nafn"] == "h2", c["leyfi"]] = .1
        with self.assertRaisesRegex(ValueError, "h2"):
            data.reikna_virkt_hlutfall(t, c, ["leyfi"])

    def test_invalid_deductions(self):
        c = data.velja_manadarblokk(self.t, 11)
        for items in (None, "leyfi", ["leyfi", "leyfi"], ["hlutfall"], ["unknown"]):
            with self.assertRaises(ValueError):
                data.reikna_virkt_hlutfall(self.t, c, items)

    def test_shifts(self):
        self.assertEqual(data.hreinsa_vaktir(self.t)["h1"], ["MV", "KV", "NV"])
        for x in (None, "12-20-KV", "MV-KV-4", "MV--KV"):
            t = self.t.copy()
            t.loc[t["Nafn"] == "h1", "Vaktir"] = x
            with self.assertRaisesRegex(ValueError, "h1"):
                data.hreinsa_vaktir(t)

    def test_skills(self):
        self.assertIn("A", data.hreinsa_haefni(self.t)["h1"])
        for x in (11, None, "100%"):
            t = self.t.copy()
            t.loc[t["Nafn"] == "h1", "a"] = x
            with self.assertRaisesRegex(ValueError, "h1"):
                data.hreinsa_haefni(t)

    def test_unrounded_target(self):
        self.assertEqual(data.reikna_markvaktir({"h1": .525}, 2026, 11, 20)["h1"], 10.5)
        for full in (None, 0, -1, float("nan"), "20%"):
            with self.assertRaises(ValueError):
                data.reikna_markvaktir({}, 2026, 11, full)

    def test_demand(self):
        q = data.lesa_og_hreinsa_monnun(self.folder)
        self.assertEqual(q, {("MV", "A"): 2, ("KV", "A"): 2, ("NV", "A"): 1})
        skrifa_daemi(self.folder, unknown_role=True)
        with self.assertRaisesRegex(ValueError, "Óþekkt"):
            data.lesa_og_hreinsa_monnun(self.folder)

    def test_pipeline_and_zero(self):
        g = data.undirbua_gogn(2026, 11, self.folder, **self.settings)
        self.assertEqual(g["starfsmenn"], ["h1"])
        self.assertEqual(g["mark"]["h1"], 10.5)
        self.assertEqual(sum(g["monnunar_thorf"].values()), 150)
        data.sannreyna_gogn(g)

    def test_ids_and_demand_checks(self):
        original = data.undirbua_gogn(2026, 11, self.folder, **self.settings)
        for mutation in ("id", "day", "fraction", "no_skill", "missing_pair"):
            g = copy.deepcopy(original)
            if mutation == "id":
                g["mark"]["h9"] = 1
            elif mutation == "day":
                g["dagar"].pop()
            elif mutation == "fraction":
                g["monnunar_thorf"][next(iter(g["monnunar_thorf"]))] = 1.5
            elif mutation == "no_skill":
                g["haefni"]["h1"] = ["V"]
            else:
                del g["monnunar_thorf"][next(iter(g["monnunar_thorf"]))]
            with self.assertRaises(ValueError):
                data.sannreyna_gogn(g)

    def test_overrides_and_stale(self):
        c = data.velja_manadarblokk(self.t, 11)
        l = {"starfsmadur": "h1", "dalkur": "Vaktir", "gamalt": " MV - KV - NV ",
             "nytt": "NV", "astaeda": "Tilbúið prófdæmi"}
        new, log = data._beita_leidrettingum(self.t, [l], 11, c)
        self.assertEqual(data.hreinsa_vaktir(new)["h1"], ["NV"])
        self.assertEqual(log[0]["gamalt"], l["gamalt"])
        self.assertEqual(data.hreinsa_vaktir(self.t)["h1"], ["MV", "KV", "NV"])
        with self.assertRaisesRegex(ValueError, "Úrelt"):
            data._beita_leidrettingum(new, [l], 11, c)
        with self.assertRaises(ValueError):
            data._beita_leidrettingum(self.t, [{**l, "astaeda": ""}], 11, c)

    def test_monthly_override_scope(self):
        c = data.velja_manadarblokk(self.t, 11)
        l = {"starfsmadur": "h1", "dalkur": "hlutfall", "manudur": 11,
             "gamalt": .525, "nytt": .7, "astaeda": "Tilbúið prófdæmi"}
        new, log = data._beita_leidrettingum(self.t, [l], 11, c)
        self.assertEqual(new.loc[new["Nafn"] == "h1", c["hlutfall"]].iloc[0], .7)
        new, log = data._beita_leidrettingum(self.t, [l], 10, data.velja_manadarblokk(self.t, 10))
        self.assertEqual(log, [])

    def test_text_requires_policy(self):
        from openpyxl import load_workbook
        path = self.folder / data.SKRA
        w = load_workbook(path)
        w["Hjúkkur"].cell(2, 2, "80% með verkefnavinnu")
        w.save(path)
        with self.assertRaisesRegex(ValueError, "texti"):
            data.undirbua_gogn(2026, 11, self.folder, **self.settings)
        g = data.undirbua_gogn(2026, 11, self.folder, textastefna="manadarprosentur", **self.settings)
        self.assertIn("h1", g["texti"])

    def test_explicit_assumptions(self):
        with self.assertRaisesRegex(ValueError, "jafnlangar"):
            data.undirbua_gogn(2026, 11, self.folder)

    def test_pipeline_explicit_override(self):
        from openpyxl import load_workbook
        path = self.folder / data.SKRA
        w = load_workbook(path)
        s = w["Hjúkkur"]
        c = [cell.value for cell in s[1]].index("Vaktir") + 1
        s.cell(2, c, "ÓÞEKKT")
        w.save(path)
        with self.assertRaisesRegex(ValueError, "h1"):
            data.undirbua_gogn(2026, 11, self.folder, **self.settings)
        l = {"starfsmadur": "h1", "dalkur": "Vaktir", "gamalt": "ÓÞEKKT",
             "nytt": "MV-KV-NV", "astaeda": "Tilbúin prófleiðrétting"}
        g = data.undirbua_gogn(2026, 11, self.folder, leidrettingar=[l], **self.settings)
        self.assertEqual(len(g["leidrettingar"]), 1)
        self.assertEqual(g["leidrettingar"][0]["gamalt"], "ÓÞEKKT")
        self.assertEqual(load_workbook(path)["Hjúkkur"].cell(2, c).value, "ÓÞEKKT")

    def test_bn_logged(self):
        skrifa_daemi(self.folder, bn=True)
        g = data.undirbua_gogn(2026, 8, self.folder, bn_sem_nm=True, **self.settings)
        self.assertEqual(g["leidrettingar"][0]["gamalt"], "BN")
        self.assertEqual(g["leidrettingar"][0]["nytt"], "NM")

    def test_unknown_shift_header(self):
        from openpyxl import load_workbook
        path = self.folder / data.SKRA
        w = load_workbook(path)
        w["Sheet2"].cell(2, 5, "XV")
        w.save(path)
        with self.assertRaisesRegex(ValueError, "vaktarfyrirsögn"):
            data.lesa_og_hreinsa_monnun(self.folder)


@unittest.skipUnless(os.getenv("HR_GAGNAMAPPA"), "HR_GAGNAMAPPA ekki skilgreind")
class RealDataTests(unittest.TestCase):
    def test_real_source(self):
        folder = Path(os.environ["HR_GAGNAMAPPA"])
        self.assertEqual(len(data.lesa_starfsfolk(folder)), 170)
        q = data.lesa_og_hreinsa_monnun(folder)
        self.assertEqual([sum(v for (ss, r), v in q.items() if ss == s)
                          for s in data.VAKTIR], [19, 19, 12])

    def test_real_november_rejects_negative(self):
        with self.assertRaisesRegex(ValueError, "h100"):
            data.undirbua_gogn(2026, 11, os.environ["HR_GAGNAMAPPA"],
                fra_drattur=["verkefni", "bor", "namsleyfi", "stjornun", "faeding", "leyfi", "veikindi"],
                vaktir_100=20, textastefna="manadarprosentur",
                jafnlangar_vaktir=True, sama_thorf_alla_daga=True)

    def test_real_october_unknown_codes_and_flag(self):
        folder = os.environ["HR_GAGNAMAPPA"]
        t = data.lesa_starfsfolk(folder)
        c = data.velja_manadarblokk(t, 10)
        t = data.breyta_prosentum(t, list(c.values()))
        ratios = data.reikna_virkt_hlutfall(t, c,
            ["verkefni", "bor", "namsleyfi", "stjornun", "faeding", "leyfi", "veikindi"])
        active = t.loc[t["Nafn"].map(ratios) > 0]
        self.assertEqual(len(active), 139)
        self.assertAlmostEqual(sum(ratios.values()), 98.579)
        with self.assertRaisesRegex(ValueError, "h74"):
            data.hreinsa_vaktir(active)
        with self.assertRaisesRegex(ValueError, "h147"):
            data.hreinsa_haefni(active)


if __name__ == "__main__":
    unittest.main()
