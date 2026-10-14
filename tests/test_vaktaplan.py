"""Prófa myndrænan útflutning án Gurobi-lausnarleitar."""
import csv
from datetime import date, timedelta
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
from greining import bua_til_vaktatoflu, velja_synishorn, vista_vaktaplan, fingrafar_gagna


class VaktaplanTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name)
        days = [date(2026, 11, 2) + timedelta(days=k) for k in range(3)]
        self.g = {"ar": 2026, "manudur": 11, "dagar": days,
                  "starfsmenn": ["h10", "h2"], "vaktir": ("MV", "KV", "NV"),
                  "leyfdar": {"h10": ["MV"], "h2": ["MV"]},
                  "haefni": {"h10": ["V"], "h2": ["V"]},
                  "mark": {"h10": 0, "h2": 1},
                  "monnunar_thorf": {(days[0], "MV", "V"): 1}}
        self.rows = [["h2", "2026-11-02", "MV", "V"]]
        self.write()
        self.run = {"solution_count": 1, "objective": 0, "gap": 0,
                    "data_fingerprint": fingrafar_gagna(self.g)}
        self.write_run()

    def tearDown(self):
        self.tmp.cleanup()

    def write(self):
        with (self.folder / "assignments.csv").open("w", newline="") as f:
            w = csv.writer(f)
            w.writerow(["nurse_id", "date", "shift", "role"])
            w.writerows(self.rows)

    def write_run(self):
        (self.folder / "run.json").write_text(json.dumps(self.run))

    def test_natural_order_and_staff_without_assignments(self):
        folk, days, cells = bua_til_vaktatoflu(self.g, self.folder / "assignments.csv")
        self.assertEqual(folk, ["h2", "h10"])
        self.assertEqual(len(days), 3)
        self.assertEqual(cells["h2", days[0]], ("MV", "V"))
        self.assertNotIn(("h10", days[0]), cells)

    def test_duplicate_assignments_are_not_hidden(self):
        self.rows.append(self.rows[0])
        self.write()
        with self.assertRaisesRegex(ValueError, "Tvær úthlutanir"):
            bua_til_vaktatoflu(self.g, self.folder / "assignments.csv")

    def test_unknown_shift_is_rejected(self):
        self.rows[0][2] = "Óþekkt"
        self.write()
        with self.assertRaises(ValueError):
            bua_til_vaktatoflu(self.g, self.folder / "assignments.csv")

    def test_selection_uses_inputs_and_fills_small_groups(self):
        self.assertEqual(velja_synishorn(self.g), ["h2", "h10"])
        self.assertEqual(velja_synishorn(self.g, 1), ["h2"])
        with self.assertRaises(ValueError):
            velja_synishorn(self.g, 0)

    def test_export_dates_roles_colours_and_empty_rows(self):
        out = self.folder / "myndir"
        info = vista_vaktaplan(self.g, self.folder, out)
        self.assertEqual(info["assignment_count"], 1)
        self.assertEqual(info["selected_nurses"], ["h2", "h10"])
        wb = load_workbook(out / "vaktaplan.xlsx")
        ws = wb.active
        self.assertEqual(ws["A7"].value, "h2")
        self.assertEqual(ws["A8"].value, "h10")
        self.assertEqual(ws["B7"].value, "MV · V")
        self.assertEqual(ws["B7"].fill.fgColor.rgb[-6:], "B8D8F0")
        self.assertIsNone(ws["B8"].value)
        self.assertEqual(ws["B8"].fill.fgColor.rgb[-6:], "FFFFFF")
        self.assertEqual(ws["B6"].value.date(), self.g["dagar"][0])
        self.assertEqual(ws.freeze_panes, "B7")
        wb.close()
        self.assertGreater((out / "vaktaplan_synishorn.png").stat().st_size, 1000)

    def test_mismatched_data_is_rejected_before_export(self):
        self.run["data_fingerprint"] = "röng gögn"
        self.write_run()
        with self.assertRaisesRegex(ValueError, "Gögnin passa ekki"):
            vista_vaktaplan(self.g, self.folder)
        self.assertFalse((self.folder / "vaktaplan.xlsx").exists())

    def test_failed_checker_prevents_export(self):
        self.rows = []
        self.write()
        with self.assertRaisesRegex(ValueError, "óháðri yfirferð"):
            vista_vaktaplan(self.g, self.folder)
        self.assertFalse((self.folder / "vaktaplan.xlsx").exists())
