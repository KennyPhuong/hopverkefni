"""Mæla gæði úr assignments.csv; engar nýjar skorður eða Gurobi-háð föll.

Frávik eru vaktareiningar, ekki staðfest launaleg yfirvinna/vanvinna.
Sanngirnismælingar lýsa úthlutunum; staðfesta ekki að lausnin sé sanngjörn.
"""
from collections import Counter, defaultdict
import csv
from datetime import date, timedelta
import hashlib
import json
import re
from pathlib import Path
from statistics import mean, pstdev

DAGUR = timedelta(days=1)
VAKTALITIR = {"MV": "B8D8F0", "KV": "F9D19A", "NV": "D3B9E8"}


def _audkennisrod(nafn):
    """Náttúruleg röð auðkenna: h2 á undan h10."""
    m = re.fullmatch(r"h(\d+)", nafn)
    return (0, int(m.group(1))) if m else (1, nafn)


def bua_til_vaktatoflu(gogn, assignments):
    """Lesa eina úthlutun á dag; varðveita líka starfsfólk án vakta."""
    folk = sorted(gogn["starfsmenn"], key=_audkennisrod)
    dagar = list(gogn["dagar"])
    reitir = {}
    with Path(assignments).open(encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames != ["nurse_id", "date", "shift", "role"]:
            raise ValueError("Ógilt dálkasnið í assignments.csv.")
        for row in reader:
            i, d = row["nurse_id"], date.fromisoformat(row["date"])
            s, r = row["shift"], row["role"]
            if (i not in folk or d not in dagar or s not in gogn["leyfdar"][i]
                    or r not in gogn["haefni"][i]
                    or gogn["monnunar_thorf"].get((d, s, r), 0) <= 0):
                raise ValueError(f"Ógild úthlutun í vaktamynd: {row}.")
            if (i, d) in reitir:
                raise ValueError(f"Tvær úthlutanir hjá {i} á {d}; ekki fela tvítekningu.")
            reitir[i, d] = (s, r)
    return folk, dagar, reitir


def velja_synishorn(gogn, hamark=15):
    """Fyrstu fimm úr þremur vaktamynstrum; fylla laus sæti eftir auðkenni.

    Val byggist á leyfðum vöktum, aldrei gæðum úthlutunarinnar.
    """
    if isinstance(hamark, bool) or not isinstance(hamark, int) or hamark < 1:
        raise ValueError("Hámark sýnishorns þarf að vera jákvæð heiltala.")
    folk = sorted(gogn["starfsmenn"], key=_audkennisrod)
    val = []
    for pattern in ({"MV", "KV", "NV"}, {"MV", "KV"}, {"MV"}):
        hopur = [i for i in folk if set(gogn["leyfdar"][i]) == pattern]
        val.extend(hopur[:5])
    val = val[:hamark]
    val.extend(i for i in folk if i not in val)
    mynstur = [("MV", "KV", "NV"), ("MV", "KV"), ("MV",)]
    def birtingarrod(i):
        pattern = tuple(s for s in ("MV", "KV", "NV") if s in gogn["leyfdar"][i])
        return (mynstur.index(pattern) if pattern in mynstur else 3, _audkennisrod(i))
    return sorted(val[:hamark], key=birtingarrod)


def _vista_excel_vaktaplan(folk, dagar, reitir, slod, heimild):
    """Heildarplan með hlutverkum; opnast án Gurobi eða annarra kerfa."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Vaktaplan"
    ws.sheet_view.showGridLines = False
    ws.cell(2, 1, f"Vaktaplan {dagar[0]:%Y-%m}").font = Font(name="Arial", size=14, bold=True)
    ws.cell(3, 1, f"Heimild: {heimild}; NV tilheyrir upphafsdegi.")
    ws.cell(4, 1, "Litaskýring")
    for col, (s, color) in zip((2, 4, 6), VAKTALITIR.items()):
        cell = ws.cell(4, col, s)
        cell.fill = PatternFill("solid", fgColor=color)
    ws.cell(4, 8, "Hvítt = frí")
    ws.cell(6, 1, "Starfsmaður")
    ws.column_dimensions["A"].width = 16
    ws.row_dimensions[6].height = 32
    vardagur = ("mán", "þri", "mið", "fim", "fös", "lau", "sun")
    for col, d in enumerate(dagar, 2):
        c = ws.cell(6, col, d)
        c.number_format = 'd" ' + vardagur[d.weekday()] + '"'
        c.fill = PatternFill("solid", fgColor="596579" if d.weekday() >= 5 else "26384D")
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = 10.5
    ws.cell(6, 1).fill = PatternFill("solid", fgColor="26384D")
    ws.cell(6, 1).font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    for row, i in enumerate(folk, 7):
        ws.cell(row, 1, i).font = Font(name="Arial", size=10)
        ws.row_dimensions[row].height = 24
        for col, d in enumerate(dagar, 2):
            entry = reitir.get((i, d))
            cell = ws.cell(row, col, f"{entry[0]} · {entry[1]}" if entry else "")
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.fill = PatternFill("solid", fgColor=VAKTALITIR[entry[0]] if entry else "FFFFFF")
            if d.weekday() == 0:
                cell.border = Border(left=Side(style="thin", color="C3CBD5"))
    ws.freeze_panes = "B7"
    ws.auto_filter.ref = f"A6:{get_column_letter(len(dagar)+1)}{len(folk)+6}"
    ws.print_title_rows = "1:6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(slod)


def _teikna_vaktaplan(gogn, dagar, reitir, val, slod):
    """Læsilegt útsnið fyrir skýrslu; aðeins vaktategund í litareitum."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle, Patch

    fig, ax = plt.subplots(figsize=(12.8, max(4, len(val)*0.31 + 1.8)))
    for row, i in enumerate(val):
        for col, d in enumerate(dagar):
            entry = reitir.get((i, d))
            color = "#" + VAKTALITIR[entry[0]] if entry else "white"
            ax.add_patch(Rectangle((col, row), 1, 1, facecolor=color,
                                   edgecolor="#E0E4EA", linewidth=.5))
            if entry:
                ax.text(col+.5, row+.5, entry[0], ha="center", va="center", fontsize=12)
    ax.set(xlim=(0, len(dagar)), ylim=(len(val), 0))
    ax.set_xticks([k+.5 for k in range(len(dagar))], [str(d.day) for d in dagar])
    labels = [i + "  " + "/".join(s for s in ("MV", "KV", "NV") if s in gogn["leyfdar"][i]) for i in val]
    ax.set_yticks([k+.5 for k in range(len(val))], labels)
    ax.xaxis.tick_top()
    ax.tick_params(length=0, labelsize=11, pad=5)
    for tick, d in zip(ax.get_xticklabels(), dagar):
        if d.weekday() >= 5:
            tick.set_bbox(dict(facecolor="#E2E7ED", edgecolor="none", pad=3))
    for spine in ax.spines.values():
        spine.set_visible(False)
    ax.set_title(f"Vaktaplan {dagar[0]:%Y-%m} — {len(val)} af {len(gogn['starfsmenn'])} starfsmönnum",
                 fontsize=15, loc="left", pad=24)
    legend = [Patch(facecolor="#"+c, label=s) for s,c in VAKTALITIR.items()]
    legend.append(Patch(facecolor="white", edgecolor="#CBD2DA", label="Frí"))
    ax.legend(handles=legend, loc="upper left", bbox_to_anchor=(0, -.065),
              ncol=4, frameon=False, fontsize=12)
    fig.text(.02, .02, "Grár dagahaus = laugardagur/sunnudagur. Hlutverk eru í Excel-heildarplaninu.", fontsize=10)
    fig.subplots_adjust(left=.15, right=.99, top=.85, bottom=.16)
    fig.savefig(slod, dpi=300, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def vista_vaktaplan(gogn, mappa, nidurstodumappa=None):
    """Búa til Excel og skýrslumynd úr yfirfarinni lausn, án nýrrar bestunar."""
    from check_solution import sannreyna_lausn
    mappa = Path(mappa)
    run = json.loads((mappa / "run.json").read_text(encoding="utf-8"))
    if not run.get("solution_count"):
        raise ValueError("Engin fundin lausn til að teikna.")
    if run.get("data_fingerprint") != fingrafar_gagna(gogn):
        raise ValueError("Gögnin passa ekki við valda keyrslu; ekki blanda lausnum og gögnum.")
    check = sannreyna_lausn(gogn, mappa / "assignments.csv")
    if not check["pass"]:
        raise ValueError("Lausn féll í óháðri yfirferð; ekki birta hana sem gilt vaktaplan.")
    out = Path(nidurstodumappa) if nidurstodumappa else mappa
    folk, dagar, reitir = bua_til_vaktatoflu(gogn, mappa / "assignments.csv")
    val = velja_synishorn(gogn)
    out.mkdir(parents=True, exist_ok=True)
    _vista_excel_vaktaplan(folk, dagar, reitir, out / "vaktaplan.xlsx", mappa.name)
    _teikna_vaktaplan(gogn, dagar, reitir, val, out / "vaktaplan_synishorn.png")
    metadata = {"source_run": mappa.name, "source_assignments_sha256": hashlib.sha256(
        (mappa / "assignments.csv").read_bytes()).hexdigest(),
        "staff_count": len(folk), "day_count": len(dagar), "assignment_count": len(reitir),
        "selected_nurses": val,
        "selection_rule": "Fyrstu fimm eftir tölulegu auðkenni úr MV/KV/NV, MV/KV og MV; laus sæti fyllt eftir auðkenni.",
        "objective": run["objective"], "gap": run["gap"]}
    (out / "vaktaplan_mynd.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    caption = (f"Útsnið úr bestu fundnu lausn valinnar keyrslu fyrir {dagar[0]:%Y-%m}. "
               f"Sýndir eru {len(val)} af {len(folk)} starfsmönnum. "
               + metadata["selection_rule"] + " Litir tákna MV, KV og NV; hvítir reitir frí. "
               "NV tilheyrir upphafsdegi. Heildarplan með hlutverkum fylgir í Excel.\n")
    (out / "vaktaplan_myndatexti.txt").write_text(caption, encoding="utf-8")
    return metadata


def skrifa_toflu(slod, radir, haus):
    with Path(slod).open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=haus)
        w.writeheader()
        w.writerows(radir)


def fingrafar_gagna(gogn):
    """Stöðugt fingrafar þeirra gagna og forsendna sem líkanið notar."""
    data = {
        "ar": gogn["ar"], "manudur": gogn["manudur"],
        "dagar": [d.isoformat() for d in gogn["dagar"]],
        "starfsmenn": [{"id": i, "mark": gogn["mark"][i],
            "leyfdar": sorted(gogn["leyfdar"][i]), "haefni": sorted(gogn["haefni"][i])}
            for i in sorted(gogn["starfsmenn"])],
        "thorf": [[d.isoformat(), s, r, q] for (d, s, r), q in sorted(gogn["monnunar_thorf"].items())],
        "saga": [[i, d.isoformat(), s] for (i, d), s in sorted(gogn.get("fyrri_vaktir", {}).items())],
        "forsendur": gogn.get("forsendur", {})}
    raw = json.dumps(data, ensure_ascii=False, sort_keys=True, allow_nan=False).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def meta_lausn(gogn, assignments, helgarhopar):
    """Skila samantekt, starfsmannatöflu, vaktatöflu og sambærilegum NV-hópum.

    Kallandi keyrir óháðan checker fyrst. Hér eru eingöngu mældar gildar
    úthlutanir með einn viðmiðunarhelgarhóp (0/1/2) fyrir hvern starfsmann.
    """
    if set(helgarhopar) != set(gogn["starfsmenn"]) or any(g not in (0, 1, 2) for g in helgarhopar.values()):
        raise ValueError("Vantar gildan viðmiðunarhelgarhóp fyrir alla starfsmenn.")
    shifts, roles, seats = Counter(), Counter(), Counter()
    helgardagar, auka, helgar = Counter(), Counter(), defaultdict(set)
    seen = set()
    with Path(assignments).open(encoding="utf-8", newline="") as f:
        for row in csv.DictReader(f):
            i, d, s, r = row["nurse_id"], date.fromisoformat(row["date"]), row["shift"], row["role"]
            if (i not in gogn["starfsmenn"] or d not in gogn["dagar"]
                    or s not in gogn["leyfdar"][i] or r not in gogn["haefni"][i]
                    or gogn["monnunar_thorf"].get((d, s, r), 0) <= 0 or (i, d) in seen):
                raise ValueError("Gæðamat þarf gildar úthlutanir; keyrið checker fyrst.")
            seen.add((i, d))
            shifts[i, s] += 1
            roles[i, r] += 1
            seats[d, s, r] += 1
            if d.weekday() >= 5 or (d.weekday() == 4 and s in ("KV", "NV")):
                fos = d - (d.weekday() - 4) * DAGUR
                group = (fos.toordinal() // 7) % 3
                helgardagar[i] += 1
                helgar[i].add(fos)
                auka[i] += int(group != helgarhopar[i])

    nurses = []
    for i in gogn["starfsmenn"]:
        fjoldi = sum(shifts[i, s] for s in gogn["vaktir"])
        mark = gogn["mark"][i]
        fravik = fjoldi - mark
        row = {"nurse_id": i, "target": mark, "assigned": fjoldi,
            "below_target": max(0, -fravik), "above_target": max(0, fravik),
            "deviation": fravik, "relative_deviation": fravik / mark if mark > 0 else None,
            "MV": shifts[i, "MV"], "KV": shifts[i, "KV"], "NV": shifts[i, "NV"],
            "night_share": shifts[i, "NV"] / fjoldi if fjoldi else None,
            "shift_pattern": "-".join(s for s in gogn["vaktir"] if s in gogn["leyfdar"][i]),
            "weekend_days": helgardagar[i], "weekends_worked": len(helgar[i]),
            "reference_group": helgarhopar[i], "extra_weekend_days": auka[i],
            "V_eligible": "V" in gogn["haefni"][i], "V_count": roles[i, "V"],
            "V_share": roles[i, "V"] / fjoldi if fjoldi else None,
            "T_eligible": "T" in gogn["haefni"][i], "T_count": roles[i, "T"],
            "T_share": roles[i, "T"] / fjoldi if fjoldi else None}
        nurses.append(row)

    coverage = []
    for s in gogn["vaktir"]:
        need = sum(q for (d, ss, r), q in gogn["monnunar_thorf"].items() if ss == s)
        assigned = sum(n[s] for n in nurses)
        missing = sum(max(0, q - seats[d, ss, r]) for (d, ss, r), q in gogn["monnunar_thorf"].items() if ss == s)
        surplus = sum(max(0, seats[d, ss, r] - q) for (d, ss, r), q in gogn["monnunar_thorf"].items() if ss == s)
        coverage.append({"shift": s, "required": need, "assigned": assigned,
                         "missing": missing, "overstaffing": surplus})

    # Aðeins fólk með NV-heimild; jafn vaktamark og sama vaktamynstur í hverjum hópi.
    # Þetta tekur ekki tillit til óstaðfestra NV-óska í frjálsum texta.
    groups = defaultdict(list)
    for n in nurses:
        if "NV" in gogn["leyfdar"][n["nurse_id"]]:
            groups[n["shift_pattern"], round(n["target"], 6)].append(n)
    night_groups = []
    for (pattern, target), group in sorted(groups.items()):
        nights = [n["NV"] for n in group]
        night_groups.append({"shift_pattern": pattern, "target": target, "nurses": len(group),
            "min_NV": min(nights), "max_NV": max(nights), "mean_NV": mean(nights),
            "range_NV": max(nights) - min(nights),
            "nurse_ids": ";".join(n["nurse_id"] for n in group)})
    deviations = [n["deviation"] for n in nurses]
    relative = [n["relative_deviation"] for n in nurses if n["relative_deviation"] is not None]
    worst = max(nurses, key=lambda n: abs(n["deviation"]), default=None)
    worst_relative = max((n for n in nurses if n["relative_deviation"] is not None),
                         key=lambda n: abs(n["relative_deviation"]), default=None)
    summary = {"below_target": sum(n["below_target"] for n in nurses),
        "above_target": sum(n["above_target"] for n in nurses),
        "total_workload_deviation": sum(abs(v) for v in deviations),
        "mean_deviation": mean(deviations) if deviations else None,
        "std_deviation": pstdev(deviations) if deviations else None,
        "std_relative_deviation": pstdev(relative) if relative else None,
        "max_absolute_deviation": abs(worst["deviation"]) if worst else None,
        "worst_absolute_nurse": worst["nurse_id"] if worst else None,
        "max_absolute_relative_deviation": abs(worst_relative["relative_deviation"]) if worst_relative else None,
        "worst_relative_nurse": worst_relative["nurse_id"] if worst_relative else None,
        "extra_weekend_days": sum(auka.values()),
        "max_extra_weekend_days": max(auka.values(), default=0),
        "overstaffing": sum(c["overstaffing"] for c in coverage),
        "unfilled_slots": sum(c["missing"] for c in coverage),
        "assignment_count": len(seen), "night_count": sum(n["NV"] for n in nurses)}
    return summary, nurses, coverage, night_groups


def vista_gaedamat(gogn, mappa, helgarhopar):
    summary, nurses, coverage, night_groups = meta_lausn(gogn, mappa / "assignments.csv", helgarhopar)
    skrifa_toflu(mappa / "quality_by_nurse.csv", nurses, list(nurses[0]) if nurses else ["nurse_id"])
    skrifa_toflu(mappa / "quality_by_shift.csv", coverage, list(coverage[0]))
    skrifa_toflu(mappa / "night_groups.csv", night_groups,
        ["shift_pattern", "target", "nurses", "min_NV", "max_NV", "mean_NV", "range_NV", "nurse_ids"])
    (mappa / "quality.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    return summary


def teikna_keyrslu(mappa):
    """Vista framvindu og stærstu vinnuskyldufrávik sem sjálfstæð PNG-myndrit."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    with (mappa / "progress.csv").open(encoding="utf-8") as f:
        progress = list(csv.DictReader(f))
    fig, axes = plt.subplots(1, 2, figsize=(11, 4), layout="constrained")
    for name, label in (("incumbent", "Besta fundna lausn"), ("best_bound", "Neðri mörk")):
        data = [(float(r["elapsed_seconds"]), float(r[name])) for r in progress if r[name]]
        if data:
            axes[0].plot(*zip(*data), label=label)
    axes[0].set(xlabel="Tími (sek.)", ylabel="Markfallseiningar", title="Incumbent og BestBd")
    axes[0].legend()
    gaps = [(float(r["elapsed_seconds"]), float(r["gap"]) * 100) for r in progress if r["gap"]]
    if gaps:
        axes[1].plot(*zip(*gaps))
    axes[1].set(xlabel="Tími (sek.)", ylabel="Gap (%)", title="Óvissa í lausnarleit")
    fig.savefig(mappa / "solver_progress.png", dpi=180)
    plt.close(fig)
    with (mappa / "quality_by_nurse.csv").open(encoding="utf-8") as f:
        nurses = list(csv.DictReader(f))
    worst = sorted(nurses, key=lambda n: abs(float(n["deviation"])), reverse=True)[:20]
    fig, ax = plt.subplots(figsize=(10, 5), layout="constrained")
    ax.bar([n["nurse_id"] for n in worst], [float(n["deviation"]) for n in worst])
    ax.axhline(0, color="black", linewidth=.8)
    ax.set(xlabel="Starfsmaður", ylabel="Vaktir yfir (+) / undir (−) marki",
           title="20 stærstu einstaklingsfrávikin frá vinnuskyldumarki")
    ax.tick_params(axis="x", rotation=45)
    fig.savefig(mappa / "workload_deviation.png", dpi=180)
    plt.close(fig)


def teikna_naemni(mappa, radir):
    """Bera saman hráa mælikvarða, ekki mismunandi vigtuð markfallsgildi."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    valid = [r for r in radir if r.get("checker_pass")]
    if not valid:
        return
    fig, axes = plt.subplots(1, 3, figsize=(12, 4), layout="constrained")
    for ax, key, title in zip(axes, ("total_workload_deviation", "overstaffing", "extra_weekend_days"),
                             ("Heildarvinnuskyldufrávik", "Umframúthlutanir", "Aukahelgardagar")):
        ax.bar([str(r["weekend_weight"]) for r in valid], [r[key] for r in valid])
        ax.set(xlabel="Helgarvigt", ylabel="Vaktir / dagar", title=title)
    fig.suptitle("Sami mánuður og sömu tímamörk; besta lausn hverrar keyrslu, ekki staðfest optimum")
    fig.savefig(mappa / "sensitivity.png", dpi=180)
    plt.close(fig)


if __name__ == "__main__":
    import argparse
    from data import lesa_novembergogn
    parser = argparse.ArgumentParser(description="Teikna varðveitta nóvemberlausn án nýrrar Gurobi-keyrslu.")
    parser.add_argument("keyrsla", type=Path)
    parser.add_argument("--ut", type=Path, default=None, help="Önnur úttaksmappa; sjálfgefið valin keyrslumappa.")
    args = parser.parse_args()
    vista_vaktaplan(lesa_novembergogn(Path(__file__).resolve().parent), args.keyrsla, args.ut)
    print("Excel og skýrslumynd vistuð í", args.ut or args.keyrsla)
